"""Tests for the visual validation loop added to generate_almanac.py."""
import sys, os
sys.path.insert(0, os.path.dirname(os.path.dirname(__file__)))

from generate_almanac import all_rules_pass

def test_all_rules_pass_when_all_pass():
    results = {"rules": [
        {"id": 1, "status": "PASS"},
        {"id": 2, "status": "PASS"},
    ]}
    assert all_rules_pass(results) is True

def test_all_rules_pass_returns_false_on_any_fail():
    results = {"rules": [
        {"id": 1, "status": "PASS"},
        {"id": 2, "status": "FAIL"},
    ]}
    assert all_rules_pass(results) is False

def test_all_rules_pass_treats_uncertain_as_passing():
    results = {"rules": [
        {"id": 1, "status": "PASS"},
        {"id": 7, "status": "UNCERTAIN"},
    ]}
    assert all_rules_pass(results) is True

def test_all_rules_pass_handles_empty():
    assert all_rules_pass({"rules": []}) is True

from unittest.mock import patch, MagicMock
from generate_almanac import find_libreoffice, convert_to_image
import os

def test_find_libreoffice_returns_known_path():
    path = find_libreoffice()
    assert path.endswith("soffice.exe")
    assert os.path.exists(path)

def test_convert_to_image_returns_png_paths(tmp_path):
    fake_docx = tmp_path / "test.docx"
    fake_docx.write_bytes(b"fake")
    fake_pdf = tmp_path / "test.pdf"
    fake_pdf.write_bytes(b"fake")

    mock_page = MagicMock()
    mock_pix = MagicMock()
    mock_pix.save = MagicMock()
    mock_page.get_pixmap.return_value = mock_pix

    mock_fitz_doc = MagicMock()
    mock_fitz_doc.__iter__ = MagicMock(return_value=iter([mock_page]))
    mock_fitz_doc.__len__ = MagicMock(return_value=1)

    mock_word_app = MagicMock()
    mock_word_doc = MagicMock()
    mock_word_app.Documents.Open.return_value = mock_word_doc

    with patch("win32com.client.Dispatch", return_value=mock_word_app), \
         patch("generate_almanac.fitz.open", return_value=mock_fitz_doc), \
         patch("os.path.exists", return_value=True):
        result = convert_to_image(str(fake_docx), str(tmp_path))

    mock_word_doc.ExportAsFixedFormat.assert_called_once()
    mock_word_doc.Close.assert_called_once()
    assert len(result) == 1
    assert result[0].endswith(".png")

from generate_almanac import validate_image
import json

def _make_popen_mock(output_text, returncode=0):
    """Helper: mock subprocess.Popen to return output_text from stdout."""
    mock_proc = MagicMock()
    mock_proc.stdout = iter(output_text.splitlines(keepends=True))
    mock_proc.stderr = iter([])
    mock_proc.returncode = returncode
    mock_proc.wait = MagicMock(return_value=returncode)
    return mock_proc

def test_validate_image_parses_json_from_claude(tmp_path):
    fake_img = tmp_path / "page_1.png"
    fake_img.write_bytes(b"fake")
    checklist = tmp_path / "checklist.md"
    checklist.write_text("## Rule 1: test rule\n- PASS condition: visible\n")

    expected = {"rules": [{"id": 1, "name": "test rule", "status": "PASS", "reason": "visible"}]}
    output_text = json.dumps(expected)

    with patch("generate_almanac.subprocess.Popen", return_value=_make_popen_mock(output_text)):
        result = validate_image([str(fake_img)], str(checklist))

    assert result["rules"][0]["status"] == "PASS"

def test_validate_image_strips_markdown_fences(tmp_path):
    fake_img = tmp_path / "page_1.png"
    fake_img.write_bytes(b"fake")
    checklist = tmp_path / "checklist.md"
    checklist.write_text("rule")

    inner = {"rules": [{"id": 1, "status": "PASS", "name": "r", "reason": "ok"}]}
    output_text = f"Here is my analysis.\n```json\n{json.dumps(inner)}\n```"

    with patch("generate_almanac.subprocess.Popen", return_value=_make_popen_mock(output_text)):
        result = validate_image([str(fake_img)], str(checklist))

    assert result["rules"][0]["id"] == 1

from generate_almanac import write_report

def test_write_report_creates_file(tmp_path):
    results = {"rules": [
        {"id": 1, "name": "Subchores indented", "status": "PASS", "reason": "visible nesting"},
        {"id": 3, "name": "Dates stacked", "status": "FAIL", "reason": "dates still horizontal"},
    ]}
    fix_log = [{"attempt": 1, "rules_attempted": [3], "fixes_applied": [
        {"rule": 3, "action": "Changed separator to newline"}
    ]}]
    report_file = str(tmp_path / "inspection_report.txt")
    write_report(results, attempt=2, max_attempts=3, fix_log=fix_log, output_file=report_file)

    content = open(report_file).read()
    assert "PASS" in content
    assert "FAIL" in content
    assert "Rule 1" in content
    assert "Rule 3" in content
    assert "Attempt: 2 of 3" in content
    assert "Changed separator to newline" in content

from generate_almanac import attempt_fixes, DEFAULT_FIX_FLAGS

def make_flags():
    return dict(DEFAULT_FIX_FLAGS)

def test_attempt_fixes_sets_rule6_flag():
    results = {"rules": [{"id": 6, "status": "FAIL", "name": "No empty rows", "reason": "rows present"}]}
    flags = make_flags()
    log = []
    attempt_fixes(results, flags, log, attempt=1)
    assert flags["rule6_no_empty_rows"] is True

def test_attempt_fixes_sets_rule3_flag():
    results = {"rules": [{"id": 3, "status": "FAIL", "name": "Dates stacked", "reason": "horizontal"}]}
    flags = make_flags()
    log = []
    attempt_fixes(results, flags, log, attempt=1)
    assert flags["rule3_vertical_dates"] is True

def test_attempt_fixes_records_fix_log():
    results = {"rules": [{"id": 6, "status": "FAIL", "name": "No empty rows", "reason": "rows"}]}
    flags = make_flags()
    log = []
    attempt_fixes(results, flags, log, attempt=1)
    assert len(log) == 1
    assert log[0]["attempt"] == 1
    assert any(f["rule"] == 6 for f in log[0]["fixes_applied"])

def test_attempt_fixes_does_not_touch_passing_rules():
    results = {"rules": [{"id": 3, "status": "PASS", "name": "Dates stacked", "reason": "ok"}]}
    flags = make_flags()
    log = []
    attempt_fixes(results, flags, log, attempt=1)
    assert flags["rule3_vertical_dates"] is False

from generate_almanac import generate_markdown, DEFAULT_FIX_FLAGS
import yaml

def _load_minimal_yaml():
    return yaml.safe_load("""
chores:
  - name: Test Chore
    category: [outdoor]
    starts: Jan 01
    interval: annually
    details:
      - Do the thing
lookahead: []
""")

def test_rule6_removes_empty_rows():
    data = _load_minimal_yaml()
    flags = dict(DEFAULT_FIX_FLAGS)
    flags["rule6_no_empty_rows"] = True
    md = generate_markdown(data, year=2026, fix_flags=flags)
    assert md.count("| | | | |") == 0

def test_rule6_off_keeps_empty_rows():
    data = _load_minimal_yaml()
    md = generate_markdown(data, year=2026)
    assert md.count("| | | | |") > 0

def test_rule3_stacks_dates_vertically():
    data = yaml.safe_load("""
chores:
  - name: Lawn Care
    category: [outdoor]
    schedule:
      - frequency: monthly
        months: Mar - May
lookahead: []
""")
    flags = dict(DEFAULT_FIX_FLAGS)
    flags["rule3_vertical_dates"] = True
    md = generate_markdown(data, year=2026, fix_flags=flags)
    assert "+---" in md or "+-" in md

def test_rule3_off_uses_pipe_table():
    data = _load_minimal_yaml()
    md = generate_markdown(data, year=2026)
    assert "| Chore | Category | Predicted Dates |" in md

def test_rule13_adds_bullet_to_subchore_in_grid_table():
    data = yaml.safe_load("""
chores:
  - name: House Cleaning
    category: [indoor]
    starts: Jan 01
    interval: annually
    subchores:
      - name: Dust Ceiling Fans
        details: [Wipe blades]
      - name: Vacuum Floors
        details: [Move furniture first]
lookahead: []
""")
    flags = dict(DEFAULT_FIX_FLAGS)
    flags["rule13_bullet_nonitalic"] = True
    flags["rule3_vertical_dates"] = True
    md = generate_markdown(data, year=2026, fix_flags=flags)
    assert "• Dust Ceiling Fans" in md or "•  Dust Ceiling Fans" in md

from generate_almanac import post_process_docx, add_yellow_bg_to_blockquotes

def test_post_process_calls_table_borders(tmp_path):
    from docx import Document
    fake_docx = str(tmp_path / "test.docx")
    doc = Document()
    doc.save(fake_docx)

    flags = dict(DEFAULT_FIX_FLAGS)
    with patch("generate_almanac.add_table_borders_to_docx") as mock_borders, \
         patch("generate_almanac.add_yellow_bg_to_blockquotes") as mock_yellow:
        post_process_docx(fake_docx, flags)
    mock_borders.assert_called_once_with(fake_docx)
    mock_yellow.assert_not_called()

def test_post_process_calls_yellow_bg_when_flag_set(tmp_path):
    from docx import Document
    fake_docx = str(tmp_path / "test.docx")
    doc = Document()
    doc.save(fake_docx)

    flags = dict(DEFAULT_FIX_FLAGS)
    flags["rule4_yellow_bg"] = True
    with patch("generate_almanac.add_table_borders_to_docx"), \
         patch("generate_almanac.add_yellow_bg_to_blockquotes") as mock_yellow:
        post_process_docx(fake_docx, flags)
    mock_yellow.assert_called_once_with(fake_docx)

from generate_almanac import remove_paragraph_borders

def test_remove_paragraph_borders_runs_without_error(tmp_path):
    from docx import Document
    from docx.oxml import parse_xml
    doc = Document()
    para = doc.add_paragraph("Test paragraph")
    pPr = para._element.get_or_add_pPr()
    border_xml = (
        '<w:pBdr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        '<w:bottom w:val="single" w:sz="6" w:space="1" w:color="auto"/>'
        '</w:pBdr>'
    )
    pPr.append(parse_xml(border_xml))
    docx_path = str(tmp_path / "test.docx")
    doc.save(docx_path)

    result = remove_paragraph_borders(docx_path)
    assert result is True

    doc2 = Document(docx_path)
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    borders = doc2.element.findall(".//w:pBdr", ns)
    assert len(borders) == 0

from generate_almanac import remove_empty_pages

def test_remove_empty_pages_runs_without_error(tmp_path):
    from docx import Document
    from docx.oxml import parse_xml
    doc = Document()
    doc.add_paragraph("Content page 1")
    para_break = doc.add_paragraph()
    run = para_break.add_run()
    run._element.append(parse_xml(
        '<w:br xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:type="page"/>'
    ))
    doc.add_paragraph("")
    docx_path = str(tmp_path / "test.docx")
    doc.save(docx_path)

    result = remove_empty_pages(docx_path)
    assert result is True

from generate_almanac import escalate_to_user

def test_escalate_to_user_writes_to_report(tmp_path, capsys):
    results = {"rules": [
        {"id": 1, "status": "PASS", "name": "Subchores indented", "reason": "ok"},
        {"id": 5, "status": "FAIL", "name": "No lines between sections", "reason": "lines present"},
    ]}
    fix_log = [{"attempt": 1, "rules_attempted": [5], "fixes_applied": [
        {"rule": 5, "action": "Removed paragraph borders"}
    ]}]
    report_file = str(tmp_path / "inspection_report.txt")
    with patch("builtins.input", return_value="D"):
        escalate_to_user(results, fix_log, report_file)

    captured = capsys.readouterr()
    assert "VALIDATION FAILED AFTER 3 ATTEMPTS" in captured.out
    assert "Rule 5" in captured.out

    content = open(report_file).read()
    assert "User response: D" in content
